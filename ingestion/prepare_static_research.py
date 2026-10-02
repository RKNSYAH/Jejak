"""Prepare a research workbook offline. Never connect to or write a database.

Outputs preserve every original data row, separate structurally prepared rows
from a review queue, and compile a transaction guarded by explicit hash approval.
Source accuracy/licensing still requires human review before promotion.
"""
import argparse
import hashlib
import json
import math
import re
import shutil
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CONTRACT = json.loads((Path(__file__).with_name("static_research_contract.json")).read_text(encoding="utf-8"))
VERIFIED_REPAIRS = json.loads((Path(__file__).with_name("static_research_verified_repairs.json")).read_text(encoding="utf-8"))
TABLES = CONTRACT["tables"]
METHOD = CONTRACT["method_version"]
SCENARIOS = {"living_cost_rates", "monthly_budgets"}
PROMOTABLE = {"prepared", "baseline"}
BASELINE_METHOD = METHOD + "-baseline"
STATISTICS = {name for name in TABLES if "period_start" in TABLES[name]["columns"].split()}
REGION_TYPES = {"country", "province", "regency", "city", "district", "neighborhood", "metro"}
PARENTS = {"province": {"country"}, "city": {"province"}, "regency": {"province"},
           "district": {"city", "regency"}, "neighborhood": {"district"}}
ENUMS = {
    "evidence_type": {"observed", "estimated", "derived", "unavailable"},
    "institution_type": {"university", "polytechnic", "school", "training_provider", "other"},
    "category": {"campus", "transit_stop", "station", "hospital", "public_facility", "other"},
    "osm_type": {"node", "way", "relation"}, "campus_osm_type": {"node", "way", "relation"},
    "data_scope": {"institution", "campus", "program", "unknown"},
    "metric": {"enrolled_students", "active_students", "new_student_intake", "graduates",
               "international_students", "program_enrollment"},
    "spending_tier": {"budget", "standard", "comfortable", "premium"},
}
INTEGERS = set("population male_population female_population working_age_population households urban_population rural_population labor_force employed_people unemployed_people sample_size student_count schools vocational_schools universities polytechnics training_centers public_schools private_schools hospitals public_hospitals private_hospitals health_centers clinics pharmacies hospital_beds public_transport_stops train_stations bus_stations transit_stations airport_count port_count observation_count osm_id campus_osm_id persons rent_observation_count".split())
RATES = {"unemployment_rate", "labor_force_participation_rate", "employment_percentage", "confidence", "rent_percentile"}
NUMERICS = INTEGERS | RATES | {"population_density", "road_length_km", "housing_price_index",
                              "consumer_price_index", "inflation_rate", "latitude", "longitude"}
WIDTHS = {"region_code": 64, "parent_region_code": 64, "region_name": 160,
          "institution_code": 80, "institution_name": 240, "source_name": 255,
          "code_source_name": 255, "kemendagri_code": 32, "bps_code": 32,
          "phone": 40, "kbli_2020_code": 64, "program_code": 80, "academic_year": 20,
          "housing_type": 30, "method_version": 80}


def issue(entry, code, detail, field=None, level="review"):
    value = {"code": code, "detail": detail, "level": level}
    if field:
        value["field"] = field
    entry["issues"].append(value)
    if level == "review":
        entry["status"] = "review"


def transform(entry, field, value, reason):
    old = entry["payload"].get(field)
    if old != value:
        entry["transformations"].append({"field": field, "from": old, "to": value, "reason": reason})
        entry["payload"][field] = value


def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str) and "\x00" in value:
        raise ValueError("NUL characters cannot be preserved in PostgreSQL JSON; correct the source workbook")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Non-finite workbook values cannot be preserved in JSON")
    return value


def apply_verified_repairs(entry, workbook_hash):
    """Apply only independently checked, hash-bound source corrections.

    Matching row identity and expected cells guard against applying an old
    source decision to a later workbook. Original cells remain in raw_data.
    """
    for repair in VERIFIED_REPAIRS.get(workbook_hash, []):
        if any(entry[field] != repair[field] for field in ["sheet", "row_number", "target_table"]):
            continue
        if any(entry["payload"].get(field) != value for field, value in repair["expected"].items()):
            issue(entry, "verified_repair_mismatch", "Verified repair no longer matches its expected input; do not apply it")
            continue
        for field, value in repair["changes"].items():
            transform(entry, field, value, "source_verified_repair:" + repair["id"])
        issue(entry, "source_rechecked", repair["source_finding"] + " Source: " + repair["source_url"] +
              " Checked: " + repair["checked_at"] + "; archived SHA-256: " + repair["source_snapshot_sha256"], level="warning")


def https_url(value):
    if not isinstance(value, str) or re.search(r"\s", value):
        return False
    try:
        url = urlsplit(value)
        return url.scheme == "https" and bool(url.hostname) and not url.username and not url.password
    except ValueError:
        return False


def normalize(entry):
    table = entry["target_table"]
    data = entry["payload"]
    for field, value in list(data.items()):
        if isinstance(value, str):
            if value.strip() in {"N/A", "n/a", "NA", "-"}:
                issue(entry, "placeholder", "Unknown values must be blank, not placeholders", field)
            elif not value.strip():
                transform(entry, field, None, "blank_to_null")
                value = None
        if value is None:
            continue
        if field.endswith("_at") or field in {"period_start", "period_end", "source_updated_at", "as_of"}:
            try:
                transform(entry, field, datetime.fromisoformat(str(value)).date().isoformat(), "date_only")
            except ValueError:
                issue(entry, "invalid_date", "Expected an ISO date or Excel date", field)
        elif field in {"is_active", "is_supported", "is_sample"}:
            if isinstance(value, str) and value in {"TRUE", "FALSE"}:
                transform(entry, field, value == "TRUE", "excel_boolean")
            elif not isinstance(value, bool):
                issue(entry, "invalid_boolean", "Expected TRUE, FALSE, or an allowed unknown", field)
        elif field == "osm_tags":
            try:
                tags = json.loads(value) if isinstance(value, str) else value
                if not isinstance(tags, dict):
                    raise ValueError("not an object")
                transform(entry, field, tags, "parse_osm_tags")
            except (ValueError, TypeError):
                issue(entry, "invalid_json", "OSM tags must be a JSON object", field)
        elif field == "kbli_2020_code":
            if isinstance(value, str):
                transform(entry, field, value.lower().replace(",", "_"), "canonical_kbli_group")
            if not re.fullmatch(r"[a-u](_[a-u])*", str(data[field])):
                issue(entry, "invalid_kbli", "Expected sections a-u joined with underscores", field)
        elif field in NUMERICS or field.endswith("_idr"):
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                issue(entry, "invalid_number", "Expected a finite numeric cell", field)
            elif field in INTEGERS and value != int(value):
                issue(entry, "invalid_integer", "Count/identity must be an integer", field)
            elif field in {"persons", "rent_observation_count"} and not 0 < value <= 2147483647:
                issue(entry, "invalid_positive_count", "Expected a positive PostgreSQL integer", field)
            elif field in {"osm_id", "campus_osm_id"} and value > 9007199254740991:
                issue(entry, "unsafe_osm_id", "OSM identity must be exact in JavaScript JSON", field)
            elif field in INTEGERS and value > 9223372036854775807:
                issue(entry, "integer_overflow", "Count exceeds the PostgreSQL bigint range", field)
            elif field not in {"latitude", "longitude", "inflation_rate"} and value < 0:
                issue(entry, "negative_number", "Count, money or share cannot be negative", field)
            elif field in RATES and not 0 <= value <= 1:
                issue(entry, "invalid_share", "Expected a share between 0 and 1", field)
            elif field in {"latitude", "longitude"} and not (-90 if field == "latitude" else -180) <= value <= (90 if field == "latitude" else 180):
                issue(entry, "invalid_coordinate", "Coordinate is outside WGS84 bounds", field)
            elif field == "inflation_rate" and not -1 <= value <= 10:
                issue(entry, "invalid_inflation", "Inflation is outside the database range", field)
        elif not isinstance(value, str):
            issue(entry, "invalid_text", "Codes and labels must be text cells", field)
        if field in ENUMS and data[field] not in ENUMS[field]:
            issue(entry, "invalid_enum", "Value is outside the database vocabulary", field)
        if field in WIDTHS and len(str(data[field])) > WIDTHS[field]:
            issue(entry, "column_too_long", f"Maximum length is {WIDTHS[field]}", field)

    if "osm_tags" in data and data["osm_tags"] is None:
        transform(entry, "osm_tags", {}, "unknown_tags_empty_object")
    if data.get("website") and not https_url(data["website"]):
        transform(entry, "website", None, "unverified_website_omitted_raw_preserved")
        issue(entry, "unverified_website", "Original website retained in raw_data; no HTTPS rewrite guessed", "website", "warning")
    citations = []
    if data.get("source_url"):
        citations = list(dict.fromkeys(str(data["source_url"]).splitlines()))
        if any(not https_url(url) for url in citations):
            issue(entry, "invalid_source_url", "Each source citation must be a complete HTTPS URL", "source_url")
        else:
            transform(entry, "source_url", citations[0], "primary_citation_full_list_retained")
    data["source_urls"] = citations
    if data.get("code_source_url") and not https_url(data["code_source_url"]):
        issue(entry, "invalid_code_source_url", "Official-code citation must be HTTPS", "code_source_url")
    if table in SCENARIOS:
        data["as_of"] = data.get("retrieved_at")
        data["method_version"] = METHOD
        issue(entry, "scenario_not_observation", "as_of uses the workbook retrieval date, not a common observation period; component source dates and assumptions remain in limitations", level="warning")
    for field in TABLES[table]["required"].split():
        if data.get(field) is None:
            issue(entry, "missing_required", "Required input is missing", field)
    for field, limit in [("region_code", 64), ("parent_region_code", 64), ("institution_code", 80)]:
        if data.get(field) and not re.fullmatch(r"[a-z0-9][a-z0-9_-]{0," + str(limit - 1) + r"}", str(data[field])):
            issue(entry, "invalid_code", "Stable codes must be lowercase, not display names", field)


def validate_observation(entry):
    table, data = entry["target_table"], entry["payload"]
    note = data.get("limitations") if isinstance(data.get("limitations"), str) else ""
    if table in STATISTICS:
        if not data.get("period_start") or not data.get("period_end"):
            issue(entry, "missing_observation_period", "Retrieval dates are not observation periods; source review needed")
        elif str(data["period_end"]) < str(data["period_start"]):
            issue(entry, "reversed_period", "Reporting end precedes start")
        # Some district rows were merged manually rather than marked Consolidated.
        # A shared year does not prove that registration definitions, publishers
        # and publication dates match across their independently sourced fields.
        mixed_population = table == "population" and (
            ";" in str(data.get("source_name") or "") or
            "Additional fields:" in note or "District male/female totals from" in note
        )
        if note.startswith("Consolidated ") or mixed_population:
            issue(entry, "mixed_provenance", "Split this row using the original source/period groups; do not parse prose into authoritative facts")
        if table == "wages_income" and "average_monthly_wage_idr period:" in note:
            period = re.search(r"average_monthly_wage_idr period: (\d{4}-\d{2}-\d{2}) to (\d{4}-\d{2}-\d{2})", note)
            if not period or tuple(period.groups()) != (data.get("period_start"), data.get("period_end")):
                issue(entry, "mixed_provenance", "Wage and statutory minimum describe different periods")
        measures = [v for k, v in data.items() if (k in NUMERICS or k.endswith("_idr")) and k not in {"confidence", "sample_size"} and v is not None]
        if data.get("evidence_type") == "unavailable" and measures:
            issue(entry, "unavailable_with_value", "Unavailable observations must not carry a metric value")
    if table == "population":
        totals = [data.get(field) for field in ["population", "male_population", "female_population"]]
        if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in totals) and totals[0] != sum(totals[1:]):
            issue(entry, "population_definition_mismatch", "Population and sex totals do not reconcile; review periods/definitions without changing values")
    if table == "regions" and data.get("region_type") not in REGION_TYPES:
        issue(entry, "unsupported_region_type", "Unclassified areas are not administrative regions; original row retained")
    if table == "regions" and data.get("region_type") in {"province", "city", "regency", "district", "neighborhood"} and not data.get("kemendagri_code"):
        issue(entry, "missing_official_code", "Administrative rows need a Kemendagri code")
    if table == "institutions" and data.get("is_active") is None:
        issue(entry, "unknown_activity", "Unknown activity is preserved as NULL, not defaulted to active", level="warning")
    if table == "public_places" and data.get("place_name") is None:
        issue(entry, "unnamed_place", "Raw name stays NULL; RPC uses an explicit unnamed category label", level="warning")
    if table in {"campuses", "public_places"}:
        if bool(data.get("osm_type")) != bool(data.get("osm_id")) or (data.get("osm_id") is not None and (not isinstance(data["osm_id"], (int, float)) or data["osm_id"] <= 0)):
            issue(entry, "incomplete_osm_identity", "Supply both OSM type and positive ID, or neither")
    if table == "student_enrollment":
        scope, osm_type, osm_id = data.get("data_scope"), data.get("campus_osm_type"), data.get("campus_osm_id")
        if bool(osm_type) != bool(osm_id) or (scope == "campus" and not osm_id) or (scope in {"institution", "unknown"} and osm_id):
            issue(entry, "invalid_enrollment_scope", "Campus identity must agree with the enrollment scope")
        if scope == "program" and not (data.get("program_code") or data.get("program_name")):
            issue(entry, "missing_program_identity", "Program scope needs an official code or source program name")
    if table == "housing_statistics":
        transform(entry, "evidence_type", "derived", "asking_quote_summary_is_derived")
        if isinstance(data.get("observation_count"), (int, float)) and data["observation_count"] < 5:
            issue(entry, "thin_housing_sample", "Fewer than five advertised offers; not a representative rent survey", level="warning")
        low, high = data.get("minimum_monthly_rent_idr"), data.get("maximum_monthly_rent_idr")
        if isinstance(low, (int, float)) and isinstance(high, (int, float)):
            for field in ["median_monthly_rent_idr", "average_monthly_rent_idr"]:
                if isinstance(data.get(field), (int, float)) and not low <= data[field] <= high:
                    issue(entry, "invalid_rent_range", "Summary is outside minimum/maximum", field)
    if table in SCENARIOS and data.get("evidence_type") != "estimated":
        issue(entry, "scenario_evidence_type", "Spending scenarios must remain estimated")


def identity(entry):
    data, table = entry["payload"], entry["target_table"]
    if table == "campuses" and data.get("osm_id") is None:
        return ("non_osm", data.get("institution_code"), data.get("campus_name"))
    values = dict(data)
    values["program_identity"] = ("code:" + str(data["program_code"]) if data.get("program_code") else
                                  "name:" + str(data["program_name"]) if data.get("program_name") else None)
    return tuple(values.get(field) for field in TABLES[table]["key"].split())


def include_static_baseline(entry):
    """Explicit opt-in for source-quality caveats, never invalid data or keys."""
    if entry["status"] != "review":
        return
    allowed = {"missing_observation_period", "mixed_provenance", "population_definition_mismatch", "unsupported_region_type"}
    blockers = [i for i in entry["issues"] if i["level"] == "review"]
    if any(i["code"] not in allowed for i in blockers):
        return
    data = entry["payload"]
    if any(i["code"] == "unsupported_region_type" for i in blockers):
        if not (entry["target_table"] == "regions" and data.get("region_type") == "unclassified_area" and
                data.get("is_supported") is False and data.get("kemendagri_code") is None and data.get("bps_code") is None):
            return
    if any(i["code"] == "mixed_provenance" for i in blockers):
        # A consolidated row cannot assert one observation/publication period for
        # all its fields. Field-specific dates remain verbatim in the audit/notes.
        for field in ["period_start", "period_end", "published_at", "confidence"]:
            if field in data:
                transform(entry, field, None, "baseline_no_common_period_or_confidence")
        transform(entry, "evidence_type", "estimated", "mixed_source_baseline_not_verified_observation")
    if "limitations" in data:
        warning = "Static research baseline, not independently verified. Observation periods may be unknown or field-specific; retrieval date is not an observation period. Use accepted scoped enrichment for comparable current housing/work evidence. "
        transform(entry, "limitations", warning + (data.get("limitations") or ""), "static_baseline_caveat")
    for finding in blockers:
        finding["level"] = "warning"
        finding["accepted_as_baseline"] = True
    entry["status"] = "baseline"


def validate_dependencies(entries):
    groups = defaultdict(list)
    for entry in entries:
        groups[(entry["target_table"], identity(entry))].append(entry)
    for grouped in groups.values():
        if len(grouped) > 1:
            for entry in grouped:
                issue(entry, "duplicate_import_identity", "All colliding rows held for review; none silently chosen")
    regions = {e["payload"].get("region_code"): e for e in entries if e["target_table"] == "regions"}
    institutions = {e["payload"].get("institution_code"): e for e in entries if e["target_table"] == "institutions"}
    campuses = {(e["payload"].get("institution_code"), e["payload"].get("osm_type"), e["payload"].get("osm_id")): e
                for e in entries if e["target_table"] == "campuses" and e["payload"].get("osm_id")}
    rates = {identity(e): e for e in entries if e["target_table"] == "living_cost_rates"}
    for field in ["kemendagri_code", "bps_code"]:
        codes = defaultdict(list)
        for entry in regions.values():
            if entry["payload"].get(field):
                codes[entry["payload"][field]].append(entry)
        for grouped in codes.values():
            if len(grouped) > 1:
                for entry in grouped:
                    issue(entry, "duplicate_official_code", "Official codes must identify one administrative region", field)

    def depth(entry, visiting):
        code = entry["payload"].get("region_code")
        if code in visiting:
            issue(entry, "region_cycle", "Parent hierarchy contains a cycle")
            return 0
        parent_code = entry["payload"].get("parent_region_code")
        if not parent_code:
            if entry["payload"].get("region_type") in PARENTS:
                issue(entry, "missing_parent", "Administrative region needs a parent")
            return 0
        parent = regions.get(parent_code)
        if not parent:
            issue(entry, "unknown_parent", "Parent must be supplied in this workbook")
            return 0
        result = depth(parent, visiting | {code}) + 1
        if parent["status"] not in PROMOTABLE:
            issue(entry, "parent_in_review", "Parent is not eligible for promotion")
        if entry["payload"].get("region_type") in PARENTS and parent["payload"].get("region_type") not in PARENTS[entry["payload"]["region_type"]]:
            issue(entry, "wrong_parent_level", "Parent administrative level is incompatible")
        child_code, official_parent = entry["payload"].get("kemendagri_code"), parent["payload"].get("kemendagri_code")
        if child_code and official_parent and not str(child_code).startswith(str(official_parent) + "."):
            issue(entry, "official_parent_mismatch", "Kemendagri prefix disagrees with parent")
        return result

    for entry in regions.values():
        entry["dependency_order"] = depth(entry, set())
    # Recheck dependencies until review status propagates through all levels.
    for _ in range(len(entries) + 1):
        changed = False
        for entry in entries:
            if entry["status"] not in PROMOTABLE:
                continue
            data, table = entry["payload"], entry["target_table"]
            refs = [("region_code", regions), ("institution_code", institutions)] if table != "regions" else [("parent_region_code", regions)]
            for field, catalog in refs:
                if data.get(field) and (data[field] not in catalog or catalog[data[field]]["status"] not in PROMOTABLE):
                    issue(entry, "unprepared_reference", "Referenced catalog row is missing or held for review", field)
                    changed = True
            if table == "student_enrollment" and data.get("campus_osm_id"):
                key = (data.get("institution_code"), data.get("campus_osm_type"), data.get("campus_osm_id"))
                if key not in campuses or campuses[key]["status"] not in PROMOTABLE:
                    issue(entry, "unprepared_campus", "Referenced campus is missing or held for review")
                    changed = True
            if table == "monthly_budgets":
                key = (data.get("region_code"), data.get("spending_tier"), data.get("as_of"), data.get("method_version"))
                if key not in rates or rates[key]["status"] not in PROMOTABLE:
                    issue(entry, "unprepared_scenario", "Referenced living-cost scenario is missing or held for review")
                    changed = True
            if data.get("region_code") in regions:
                level = regions[data["region_code"]]["payload"].get("region_type")
                if table in {"labor_force", "sector_employment", "wages_income", "cost_of_living"} and level in {"district", "neighborhood"}:
                    issue(entry, "unsupported_statistical_grain", "These inputs are city/province context, not kecamatan facts")
                    changed = True
        if not changed:
            break


def verify_scenarios(entries):
    rates = {identity(e): e for e in entries if e["target_table"] == "living_cost_rates"}
    for entry in entries:
        table, data = entry["target_table"], entry["payload"]
        if table not in SCENARIOS:
            continue
        def check_sum(fields, total):
            values = [data.get(f) for f in fields + [total]]
            if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in values):
                if abs(sum(values[:-1]) - values[-1]) > 0.000001:
                    issue(entry, "scenario_total_mismatch", "Cached formula/value disagrees with components", total)
        if table == "living_cost_rates":
            check_sum(["food_monthly_idr", "utilities_monthly_idr", "transport_monthly_idr", "connectivity_monthly_idr", "laundry_monthly_idr"], "living_cost_total_monthly_idr")
            for field, formula in entry["formulas"].items():
                if field != "living_cost_total_monthly_idr" or formula != f"=SUM(C{entry['row_number']}:G{entry['row_number']})":
                    issue(entry, "unsupported_formula", "Only the reviewed component SUM is permitted", field)
        else:
            check_sum(["rent_monthly_idr", "living_cost_monthly_idr"], "total_monthly_idr")
            linked = rates.get((data.get("region_code"), data.get("spending_tier"), data.get("as_of"), data.get("method_version")))
            if not linked or linked["status"] not in PROMOTABLE or data.get("living_cost_monthly_idr") != linked["payload"].get("living_cost_total_monthly_idr") or data.get("persons") != linked["payload"].get("persons"):
                issue(entry, "scenario_dependency_mismatch", "Living costs must match the same region/tier/date/person count")
            for field, formula in entry["formulas"].items():
                if field == "living_cost_monthly_idr":
                    match = re.fullmatch(r"='([^']+)'!H(\d+)", formula)
                    if not match or not linked or match.group(1) != linked["sheet"] or int(match.group(2)) != linked["row_number"]:
                        issue(entry, "scenario_formula_reference", "Formula must reference the matching region and spending tier", field)
                elif field != "total_monthly_idr" or formula != f"=D{entry['row_number']}+E{entry['row_number']}":
                    issue(entry, "unsupported_formula", "Only the reviewed rent-plus-cost formula is permitted", field)


def prepare(workbook_path, include_review_as_baseline=False):
    path = Path(workbook_path)
    workbook_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    formulas = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=False)
    entries, present, formula_count = [], set(), 0
    try:
        for sheet in formulas:
            table = re.sub(r"^\d+_", "", sheet.title)
            if table == "quality_audit":
                continue
            if table not in TABLES or table in present:
                raise ValueError(f"Unknown or repeated data sheet: {sheet.title}")
            present.add(table)
            headers = [cell.value for cell in sheet[1]]
            if any(not isinstance(h, str) for h in headers) or len(set(headers)) != len(headers):
                raise ValueError(f"Invalid/repeated headers on {sheet.title}")
            columns = TABLES[table]["columns"].split()
            unknown = set(headers) - set(columns)
            missing = set(TABLES[table]["required"].split()) - set(headers) - {"as_of", "method_version"}
            if unknown or missing:
                raise ValueError(f"{sheet.title}: unknown headers {sorted(unknown)}, missing required headers {sorted(missing)}")
            for row in sheet.iter_rows(min_row=2):
                if not any(cell.value is not None for cell in row):
                    continue
                row_number = row[0].row
                values = {header: json_value(cached[sheet.title].cell(row_number, index + 1).value) for index, header in enumerate(headers)}
                raw = {header: json_value(row[index].value) for index, header in enumerate(headers)}
                entry = {"sheet": sheet.title, "row_number": row_number, "target_table": table,
                         "status": "prepared", "dependency_order": 0, "raw_data": raw,
                         "payload": {c: values.get(c) for c in columns}, "formulas": {},
                         "issues": [], "transformations": []}
                for index, header in enumerate(headers):
                    cell = row[index]
                    if cell.data_type == "f":
                        formula_count += 1
                        entry["formulas"][header] = cell.value
                        if values[header] is None:
                            issue(entry, "missing_formula_cache", "Recalculate a separate copy in Excel/LibreOffice before retrying", header)
                        if table not in SCENARIOS:
                            issue(entry, "unsupported_formula", "Only reviewed scenario formulas are evaluated", header)
                    if cell.data_type == "e" or cached[sheet.title].cell(row_number, index + 1).data_type == "e":
                        issue(entry, "excel_error", "Source formula contains an Excel error", header)
                apply_verified_repairs(entry, workbook_hash)
                normalize(entry)
                validate_observation(entry)
                if include_review_as_baseline:
                    if table in SCENARIOS:
                        entry["payload"]["method_version"] = BASELINE_METHOD
                    include_static_baseline(entry)
                entries.append(entry)
        verify_scenarios(entries)
        validate_dependencies(entries)
    finally:
        formulas.close()
        cached.close()
    sheets = {table: dict(Counter(e["status"] for e in entries if e["target_table"] == table)) for table in TABLES if table in present}
    counts = Counter(e["status"] for e in entries)
    report = {"workbook_sha256": workbook_hash, "source_file": path.name,
              "method_version": BASELINE_METHOD if include_review_as_baseline else METHOD,
              "include_review_as_baseline": include_review_as_baseline, "baseline_rows": counts["baseline"],
              "import_rows": counts["prepared"] + counts["baseline"],
              "source_accuracy_verified": False, "input_rows": len(entries), "prepared_rows": counts["prepared"],
              "review_rows": counts["review"], "scenario_formulas_checked": formula_count,
              "sheets": sheets, "missing_data_sheets": sorted(set(TABLES) - present),
              "missing_reference_sheets": ["sector_mapping", "geospatial_sources", "estimation_parameters"],
              "issue_counts": dict(Counter(i["code"] for e in entries for i in e["issues"])),
              "transformation_counts": dict(Counter(t["reason"] for e in entries for t in e["transformations"]))}
    return {"report": report, "rows": entries}


def sql_string(value):
    return "'" + str(value).replace("'", "''") + "'"


def scope_manifest(manifest, tables):
    """Validate against the whole workbook first, then compile only this scope."""
    selected = set(tables)
    if not selected or not selected <= set(TABLES):
        raise ValueError("Scope must contain known target tables")
    rows = [row for row in manifest["rows"] if row["target_table"] in selected]
    counts = Counter(row["status"] for row in rows)
    report = dict(manifest["report"])
    report.update(scope=sorted(selected), full_workbook_input_rows=report["input_rows"],
                  input_rows=len(rows), import_rows=counts["prepared"] + counts["baseline"],
                  prepared_rows=counts["prepared"], baseline_rows=counts["baseline"], review_rows=counts["review"],
                  sheets={t: report["sheets"][t] for t in selected if t in report["sheets"]},
                  missing_data_sheets=sorted(selected - set(report["sheets"])),
                  scenario_formulas_checked=sum(len(row["formulas"]) for row in rows),
                  issue_counts=dict(Counter(i["code"] for row in rows for i in row["issues"])),
                  transformation_counts=dict(Counter(t["reason"] for row in rows for t in row["transformations"])))
    return {"report": report, "rows": rows}


def json_sql(value):
    return sql_string(json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))) + "::jsonb"


def compile_import(manifest):
    """One transaction; no connections, credentials, server files, or OS commands."""
    report, rows = manifest["report"], manifest["rows"]
    batch, method = sql_string(report["workbook_sha256"]), sql_string(report["method_version"])
    predicate = f"workbook_sha256 = {batch} and method_version = {method}"
    stage_predicate = f"s.workbook_sha256 = {batch} and s.method_version = {method}"
    statements = ["-- Generated offline. Review this artifact and the report before approving its hash.",
                  "begin;", "set local standard_conforming_strings = on;",
                  "set local search_path = pg_catalog;", "set local lock_timeout = '10s';",
                  "set local statement_timeout = '120s';",
                  f"do $$ begin if current_setting('jejak.static_import_approved_sha256', true) is distinct from {batch} then raise exception 'Static import requires explicit workbook hash approval'; end if; end $$;",
                  "select pg_catalog.pg_advisory_xact_lock(pg_catalog.hashtextextended('jejak:static-import', 0));",
                  f"insert into private.static_import_batches(workbook_sha256, method_version, source_file, report) values ({batch}, {method}, {sql_string(report['source_file'])}, {json_sql(report)}) on conflict (workbook_sha256, method_version) do update set report = excluded.report, source_file = excluded.source_file;"]
    stage_columns = "sheet row_number target_table status dependency_order raw_data payload formulas issues transformations".split()
    for start in range(0, len(rows), 250):
        chunk = [{key: e[key] for key in stage_columns} for e in rows[start:start + 250]]
        columns = ", ".join(stage_columns)
        select_columns = ", ".join("s." + c for c in stage_columns)
        statements.append(f"insert into private.static_import_rows(workbook_sha256, method_version, {columns}) select {batch}, {method}, {select_columns} from jsonb_to_recordset({json_sql(chunk)}) as s(sheet text, row_number integer, target_table text, status text, dependency_order integer, raw_data jsonb, payload jsonb, formulas jsonb, issues jsonb, transformations jsonb) on conflict (workbook_sha256, method_version, sheet, row_number) do update set " + ", ".join(f"{c} = excluded.{c}" for c in stage_columns if c not in {"sheet", "row_number"}) + ";")
    statements.append(f"do $$ begin if (select count(*) from private.static_import_rows where {predicate}) <> {len(rows)} then raise exception 'Staged row count does not reconcile'; end if; end $$;")
    # Foreign-key order is explicit. Regions are inserted one hierarchy depth at a time.
    order = ["regions", "institutions", "campuses", "public_places", "population", "labor_force", "sector_employment", "wages_income", "student_enrollment", "education_facilities", "healthcare_facilities", "transport_infrastructure", "housing_statistics", "cost_of_living", "living_cost_rates", "monthly_budgets"]
    for table in order:
        columns = TABLES[table]["columns"].split() + ["source_urls"]
        if not any(e["target_table"] == table and e["status"] in PROMOTABLE for e in rows):
            continue
        filters = [""]
        if table == "regions":
            filters = [f" and s.dependency_order = {depth}" for depth in sorted({e["dependency_order"] for e in rows if e["target_table"] == table and e["status"] in PROMOTABLE})]
        elif table == "campuses":
            filters = [" and s.payload->>'osm_id' is not null", " and s.payload->>'osm_id' is null"]
        for extra in filters:
            keys = TABLES[table]["key"].split()
            conflict_filter = ""
            if table == "campuses" and "is null" in extra:
                keys = ["institution_code", "campus_name"]
                conflict_filter = " where osm_type is null and osm_id is null"
            elif table == "public_places":
                conflict_filter = " where osm_type is not null and osm_id is not null"
            updates = [c for c in columns if c not in keys]
            assignments = []
            for column in updates:
                if table == "regions" and column in {"source_name", "source_updated_at", "source_url", "source_urls"}:
                    # The workbook does not supply boundaries. Preserve any stored
                    # geometry's provenance; the incoming catalogue is in staging.
                    assignments.append(f"{column} = case when regions.geometry is not null then regions.{column} else excluded.{column} end")
                else:
                    assignments.append(f"{column} = excluded.{column}")
            statuses = "('prepared', 'baseline')" if report.get("include_review_as_baseline") else "('prepared')"
            statements.append(f"insert into public.{table} ({', '.join(columns)}) select " + ", ".join("p." + c for c in columns) + f" from private.static_import_rows s cross join lateral jsonb_populate_record(null::public.{table}, s.payload) p where {stage_predicate} and s.status in {statuses} and s.target_table = {sql_string(table)}{extra} on conflict ({', '.join(keys)}){conflict_filter} do update set " + ", ".join(assignments) + ";")
    statements.extend([f"select target_table, status, count(*) as rows from private.static_import_rows where {predicate} group by target_table, status order by target_table, status;", "commit;"])
    return "\n".join(statements) + "\n"


def write_preparation(manifest, workbook_path, output):
    output = Path(output).resolve()
    if output == Path(workbook_path).resolve().parent or output == ROOT / "public" or ROOT / "public" in output.parents:
        raise ValueError("Preparation outputs must not be publicly served or overwrite the original folder")
    output.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(workbook_path, output / "source.xlsx")
    if hashlib.sha256((output / "source.xlsx").read_bytes()).hexdigest() != manifest["report"]["workbook_sha256"]:
        raise ValueError("Source workbook changed during preparation; rerun before using the outputs")
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8", newline="\n")
    (output / "report.json").write_text(json.dumps(manifest["report"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    (output / "review.jsonl").write_text("".join(json.dumps(e, ensure_ascii=False, allow_nan=False) + "\n" for e in manifest["rows"] if e["status"] in {"review", "baseline"}), encoding="utf-8", newline="\n")
    (output / "import.sql").write_text(compile_import(manifest), encoding="utf-8", newline="\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", nargs="?", type=Path, default=ROOT / "public/Jejak_Static_Data_Research_Completed.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "ingestion/data/prepared/static_research")
    parser.add_argument("--check-only", action="store_true", help="Print the report without creating outputs")
    parser.add_argument("--include-review-as-baseline", action="store_true", help="Include source-quality caveats as flagged static baselines; invalid values and relationships still block promotion")
    parser.add_argument("--scope", nargs="+", choices=sorted(TABLES), help="Compile selected tables after whole-workbook dependency validation")
    args = parser.parse_args()
    manifest = prepare(args.workbook, args.include_review_as_baseline)
    if args.scope:
        manifest = scope_manifest(manifest, args.scope)
    if not args.check_only:
        write_preparation(manifest, args.workbook, args.output)
    print(json.dumps(manifest["report"], ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
